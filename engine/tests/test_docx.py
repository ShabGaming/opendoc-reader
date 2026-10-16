import io
import struct
import zipfile
import zlib
from pathlib import Path

import docx
import pytest
from lxml import etree

from opendoc_engine.exporters.docx import docx_bytes
from opendoc_engine.exporters.omml import M_NS
from opendoc_engine.exporters.omml_check import check_omml
from opendoc_engine.ir import (
    Block,
    BlockType,
    Document,
    EngineInfo,
    ListItem,
    Page,
    Source,
    Table,
    TableCell,
)

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W, "m": M_NS}
BOX = (0.0, 0.0, 850.0, 400.0)


def png(path: Path, width: int = 40, height: int = 20) -> Path:
    def chunk(kind: bytes, data: bytes) -> bytes:
        body = kind + data
        return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body))

    rows = b"".join(b"\x00" + b"\xcc" * width for _ in range(height))
    path.write_bytes(
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 0, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(rows))
        + chunk(b"IEND", b"")
    )
    return path


def para(block_id: str, text: str) -> Block:
    return Block(block_id, BlockType.PARAGRAPH, BOX, 0, text=text)


def cell(row: int, col: int, text: str, **kwargs) -> TableCell:
    return TableCell(row, col, blocks=[para(f"c{row}{col}", text)], **kwargs)


def page(index: int, *blocks: Block, width: int = 1700, height: int = 2200) -> Page:
    for i, block in enumerate(blocks):
        block.order = i
    return Page(index, width, height, f"{index}.png", list(blocks))


def document(*pages: Page, path: str = "C:\\scans\\report.pdf") -> Document:
    return Document(
        Source(path, "pdf", len(pages)), EngineInfo("demo", "none", "0.1.0"), list(pages)
    )


def build(doc: Document) -> tuple[docx.document.Document, etree._Element, list[str]]:
    data, warnings = docx_bytes(doc)
    package = docx.Document(io.BytesIO(data))  # python-docx can open it
    body = etree.fromstring(zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml"))
    return package, body, warnings


def styles(package) -> list[str]:
    return [p.style.name for p in package.paragraphs]


@pytest.fixture
def sample(tmp_path) -> Document:
    crop = png(tmp_path / "p0-b9.png")
    merged = Table(
        3,
        3,
        [
            cell(0, 0, "Region", header=True, col_span=2),
            cell(0, 2, "Total", header=True),
            cell(1, 0, "North", row_span=2),
            cell(1, 1, "Q1"),
            cell(1, 2, "10"),
            cell(2, 1, "Q2"),
            cell(2, 2, "20"),
        ],
    )
    first = page(
        0,
        Block("head", BlockType.HEADER, BOX, 0, text="RUNNING HEAD"),
        Block("h1", BlockType.HEADING, BOX, 0, level=1, text="Annual report"),
        Block("h2", BlockType.HEADING, BOX, 0, level=2, text="Results"),
        para("p1", "Plain *emphasis* and **strong** text with $x^2$ inline."),
        Block("tb", BlockType.TABLE, BOX, 0, table=merged),
        Block("f", BlockType.FORMULA, BOX, 0, latex=r"\frac{a}{b} = c", number="7"),
        Block("fig", BlockType.FIGURE, (0.0, 0.0, 850.0, 400.0), 0, figure_crop=str(crop)),
        Block("cap", BlockType.CAPTION, BOX, 0, text="Figure 1. A grey box", caption_for="fig"),
        Block("fn", BlockType.FOOTNOTE, BOX, 0, text="1 A footnote."),
        para("last", "Body text after the footnote in reading order."),
        Block("pn", BlockType.PAGE_NUMBER, BOX, 0, text="1"),
    )
    second = page(
        1,
        Block("foot", BlockType.FOOTER, BOX, 0, text="Page footer"),
        Block("h3", BlockType.HEADING, BOX, 0, level=3, text="Second page"),
        para("p2", "More text."),
    )
    return document(first, second)


def test_opens_with_python_docx_and_headings_use_the_template_styles(sample):
    package, _, warnings = build(sample)
    assert warnings == []
    names = styles(package)
    assert names[:2] == ["Heading 1", "Heading 2"]
    assert "Heading 3" in names
    heading = next(p for p in package.paragraphs if p.text == "Annual report")
    assert heading.style.name == "Heading 1"


def test_template_styles_exist():
    package = docx.Document(io.BytesIO(docx_bytes(document())[0]))
    available = {s.name for s in package.styles}
    for name in ("Normal", "Heading 1", "Heading 2", "Heading 3", "Caption", "Table Grid"):
        assert name in available


def test_page_furniture_is_omitted(sample):
    package, body, _ = build(sample)
    text = "\n".join(p.text for p in package.paragraphs)
    assert "RUNNING HEAD" not in text
    assert "Page footer" not in text
    assert body.find(".//w:sectPr/w:headerReference", NS) is None
    assert body.find(".//w:sectPr/w:footerReference", NS) is None
    assert package.sections[0].header.is_linked_to_previous  # no header part of its own


def test_inline_formatting_and_inline_math(sample):
    package, _, _ = build(sample)
    paragraph = next(p for p in package.paragraphs if p.text.startswith("Plain"))
    runs = {r.text: r for r in paragraph.runs}
    assert runs["emphasis"].italic is True
    assert runs["strong"].bold is True
    inline = paragraph._p.findall("m:oMath", NS)
    assert len(inline) == 1
    assert check_omml(inline[0]) == []


def test_merged_cell_table_with_header_row(sample):
    package, _, _ = build(sample)
    assert len(package.tables) == 2  # the data table and the equation-number layout table
    table = package.tables[0]
    first_row = table.rows[0]._tr
    assert first_row.find("w:trPr/w:tblHeader", NS) is not None
    spans = [tc.find("w:tcPr/w:gridSpan", NS) for tc in first_row.findall("w:tc", NS)]
    assert spans[0].get(f"{{{W}}}val") == "2"
    # "North" spans rows 1 and 2: restart in the first, continue in the second.
    restart = table.rows[1]._tr.findall("w:tc", NS)[0].find("w:tcPr/w:vMerge", NS)
    assert restart.get(f"{{{W}}}val") == "restart"
    carry = table.rows[2]._tr.findall("w:tc", NS)[0].find("w:tcPr/w:vMerge", NS)
    assert carry is not None
    assert table.cell(1, 0).text == "North"
    assert table.cell(1, 2).text == "10"
    # header text is bold
    assert all(r.bold for p in table.cell(0, 0).paragraphs for r in p.runs)


def test_numbered_display_equation_is_native_math_with_a_right_aligned_number(sample):
    package, _, _ = build(sample)
    layout = package.tables[1]
    math_cell, number_cell = layout.rows[0].cells
    display = math_cell._tc.findall(".//m:oMathPara", NS)
    assert len(display) == 1
    assert check_omml(display[0]) == []
    assert display[0].find(".//m:f", NS) is not None
    assert number_cell.text == "(7)"
    assert number_cell.paragraphs[0].alignment == docx.enum.text.WD_ALIGN_PARAGRAPH.RIGHT
    # a layout table must not draw borders
    assert layout._tbl.tblPr.find("w:tblBorders", NS) is None
    assert layout.style.name == "Normal Table"


def test_unnumbered_display_equation_is_its_own_paragraph():
    doc = document(page(0, Block("f", BlockType.FORMULA, BOX, 0, latex=r"x^2")))
    package, _, _ = build(doc)
    assert len(package.tables) == 0
    paragraph = package.paragraphs[0]
    assert paragraph.style.name == "Equation"
    assert len(paragraph._p.findall("m:oMathPara/m:oMath", NS)) == 1


def test_all_math_in_the_document_is_valid_and_in_the_math_namespace(sample):
    _, body, _ = build(sample)
    math = body.findall(".//m:oMathPara", NS) + body.findall(".//w:p/m:oMath", NS)
    assert len(math) == 2
    for node in math:
        assert check_omml(node) == []
    for element in body.iter():
        if element.tag.startswith("{") and "officeDocument/2006/math" in element.tag:
            assert element.tag.startswith(f"{{{M_NS}}}")


def test_figure_is_inserted_and_kept_with_its_caption(sample):
    package, _, _ = build(sample)
    assert len(package.inline_shapes) == 1
    shape = package.inline_shapes[0]
    assert 0 < shape.width <= package.sections[0].page_width - 2 * package.sections[0].left_margin
    figure = next(p for p in package.paragraphs if p._p.findall(".//w:drawing", NS))
    assert figure.paragraph_format.keep_with_next is True
    caption = next(p for p in package.paragraphs if p.text.startswith("Figure 1"))
    assert caption.style.name == "Caption"


def test_page_break_before_the_second_source_page(sample):
    package, _, _ = build(sample)
    breaks = [p for p in package.paragraphs if p.paragraph_format.page_break_before]
    assert [p.text for p in breaks] == ["Second page"]


def test_no_page_break_on_the_first_page_or_between_empty_pages():
    doc = document(
        page(0, para("a", "one")),
        page(1, Block("f", BlockType.FOOTER, BOX, 0, text="only furniture")),
        page(2, para("b", "three")),
    )
    package, _, _ = build(doc)
    flags = [p.paragraph_format.page_break_before for p in package.paragraphs]
    assert flags == [None, True]


def test_page_break_before_a_table_uses_a_thin_paragraph():
    table = Table(1, 1, [cell(0, 0, "x")])
    doc = document(
        page(0, para("a", "one")), page(1, Block("t", BlockType.TABLE, BOX, 0, table=table))
    )
    package, body, _ = build(doc)
    spacer = package.paragraphs[1]
    assert spacer.text == ""
    assert spacer.paragraph_format.page_break_before is True
    children = [etree.QName(c).localname for c in body.find("w:body", NS)]
    assert children[:4] == ["p", "p", "tbl", "p"]  # text, spacer, table, closing paragraph


def test_footnotes_move_to_the_end_of_their_page(sample):
    package, _, _ = build(sample)
    texts = [p.text for p in package.paragraphs]
    note = texts.index("1 A footnote.")
    after = texts.index("Body text after the footnote in reading order.")
    assert after < note < texts.index("Second page")
    assert package.paragraphs[note].style.name == "Footnote Text"


def test_lists_are_numbered_per_list_and_nested():
    first = Block(
        "l1",
        BlockType.LIST,
        BOX,
        0,
        items=[ListItem("one", True, [ListItem("sub bullet")]), ListItem("two", True)],
    )
    second = Block("l2", BlockType.LIST, BOX, 0, items=[ListItem("again", True)])
    package, _, _ = build(document(page(0, first, second)))
    paragraphs = package.paragraphs
    assert all(p.style.name == "List Paragraph" for p in paragraphs)

    def numbering(p):
        num = p._p.find("w:pPr/w:numPr", NS)
        return int(num.find("w:numId", NS).get(f"{{{W}}}val")), int(
            num.find("w:ilvl", NS).get(f"{{{W}}}val")
        )

    (one, _), (sub, sub_level), (two, _), (again, _) = (numbering(p) for p in paragraphs)
    assert one == two != again  # a second list gets its own instance, so it restarts at 1
    assert sub != one and sub_level == 1  # unordered child uses the bullet definition
    numbering_xml = package.part.numbering_part.element
    override = numbering_xml.xpath(f"//w:num[@w:numId='{again}']/w:lvlOverride/w:startOverride")
    assert override


def test_hyperlinks_only_for_safe_schemes():
    doc = document(
        page(0, para("p", "[ok](https://example.com) and [bad](javascript:alert(1)) here"))
    )
    package, body, _ = build(doc)
    links = body.findall(".//w:hyperlink", NS)
    assert len(links) == 1
    rid = links[0].get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
    assert package.part.rels[rid].target_ref == "https://example.com"
    assert "bad" in package.paragraphs[0].text


def test_unsupported_formula_is_kept_as_latex_text_with_a_warning():
    doc = document(page(0, Block("p1-b4", BlockType.FORMULA, BOX, 0, latex=r"\cancel{x}")))
    package, body, warnings = build(doc)
    assert warnings and "formula p1-b4 kept as LaTeX text" in warnings[0]
    assert body.find(".//m:oMath", NS) is None
    paragraph = package.paragraphs[0]
    assert paragraph.text == r"\cancel{x}"
    assert paragraph.runs[0].style.name == "Code"


def test_unsupported_inline_math_is_kept_as_latex_text():
    doc = document(page(0, para("p3", r"before $\cancel{x}$ after")))
    package, _, warnings = build(doc)
    assert warnings == [warnings[0]] and "inline math in p3" in warnings[0]
    assert package.paragraphs[0].text == r"before $\cancel{x}$ after"


def test_missing_figure_crop_is_skipped_with_a_warning(tmp_path):
    block = Block("fig", BlockType.FIGURE, BOX, 0, figure_crop=str(tmp_path / "nope.png"))
    package, _, warnings = build(document(page(0, block, para("p", "text"))))
    assert len(package.inline_shapes) == 0
    assert any("figure fig skipped" in w for w in warnings)


def test_core_properties():
    doc = document(page(0, Block("h", BlockType.HEADING, BOX, 0, level=1, text="The *Title*")))
    package, _, _ = build(doc)
    props = package.core_properties
    assert props.title == "The Title"
    assert props.author == "OpenDoc Reader"
    assert props.last_modified_by == ""
    assert props.comments == ""


def test_title_falls_back_to_the_file_name():
    package, _, _ = build(document(page(0, para("p", "text")), path="C:\\scans\\Invoice 2024.pdf"))
    assert package.core_properties.title == "Invoice 2024"


def test_page_size_follows_the_first_page():
    letter = build(document(page(0, para("p", "x"), width=1700, height=2200)))[0]
    a4 = build(document(page(0, para("p", "x"), width=1654, height=2339)))[0]
    assert letter.sections[0].page_width < letter.sections[0].page_height
    assert round(letter.sections[0].page_width.inches, 1) == 8.5
    assert round(a4.sections[0].page_width.inches, 1) == 8.3


def test_empty_document_still_opens():
    package, _, warnings = build(document())
    assert package.paragraphs == []
    assert warnings == []


def test_illegal_xml_characters_are_removed():
    package, _, _ = build(document(page(0, para("p", "a\x00b\x0bc"))))
    assert package.paragraphs[0].text == "abc"


def test_package_has_no_leftover_template_parts():
    data = docx_bytes(document(page(0, para("p", "x"))))[0]
    names = zipfile.ZipFile(io.BytesIO(data)).namelist()
    assert not [n for n in names if "thumbnail" in n or "customXml" in n]
    assert "word/styles.xml" in names


def test_table_cell_blocks_keep_their_formatting_and_math():
    inner = Table(
        1,
        2,
        [
            TableCell(0, 0, blocks=[para("a", "plain *it*")]),
            TableCell(0, 1, blocks=[Block("f", BlockType.FORMULA, BOX, 0, latex="x^2")]),
        ],
    )
    package, _, _ = build(document(page(0, Block("t", BlockType.TABLE, BOX, 0, table=inner))))
    table = package.tables[0]
    assert table.cell(0, 0).paragraphs[0].runs[1].italic is True
    assert table.cell(0, 1)._tc.find(".//m:oMathPara", NS) is not None
