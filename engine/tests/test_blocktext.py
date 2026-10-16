"""Block-level Markdown must not reach IR text, the DOCX or the Markdown export.

The DeepSeek cases use answers recorded from the real model (tests/fixtures/deepseek-*.txt);
the PaddleOCR-VL and edge cases are small hand-written layout blocks in the shapes the models
are known to produce.
"""

import io
import zipfile
from pathlib import Path

import docx
import pytest
from lxml import etree

from opendoc_engine.blocktext import (
    escape_inline,
    export_text,
    literal_lead,
    normalize_blocks,
    strip_heading_markers,
)
from opendoc_engine.engines.deepseek import tidy_blocks
from opendoc_engine.exporters.docx import docx_bytes
from opendoc_engine.exporters.inline import parse_inline
from opendoc_engine.exporters.markdown import to_markdown
from opendoc_engine.grounding import parse_grounding
from opendoc_engine.ir import Block, BlockType, Document, EngineInfo, Page, Source
from opendoc_engine.paddle_blocks import LayoutItem, blocks_from_layout

FIXTURES = Path(__file__).parent / "fixtures"
BOX = (0.0, 0.0, 100.0, 20.0)
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
M = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def deepseek(name: str) -> list[Block]:
    text = (FIXTURES / f"deepseek-{name}.txt").read_text(encoding="utf-8")
    return tidy_blocks(parse_grounding(text, (1318, 2200)))


def vl(label: str, content: str) -> list[Block]:
    return blocks_from_layout([LayoutItem(label, BOX, content)], "vl")


def block(kind: BlockType, text: str, level: int | None = None) -> Block:
    return Block("b0", kind, BOX, 0, level=level, text=text)


def document(blocks: list[Block]) -> Document:
    for order, b in enumerate(blocks):
        b.order = order
    page = Page(0, 1000, 1400, "0.png", blocks)
    return Document(Source("C:\\scan.png", "image", 1), EngineInfo("x", "y", "0"), [page])


def paragraphs(doc: Document) -> tuple[list[str], list[str]]:
    data, _ = docx_bytes(doc)
    package = docx.Document(io.BytesIO(data))
    return [p.text for p in package.paragraphs], [p.style.name for p in package.paragraphs]


# --- DeepSeek, recorded -----------------------------------------------------------------


def test_deepseek_form_headings_lose_their_markers_and_keep_levels():
    headings = [b for b in deepseek("form") if b.type is BlockType.HEADING]
    assert headings
    assert not any(b.text.startswith("#") for b in headings)
    assert headings[0].text == "Request for Taxpayer Identification Number and Certification"
    assert headings[0].level == 1
    assert {b.level for b in headings[1:]} == {2}


def test_deepseek_twocol_section_title():
    heading = next(b for b in deepseek("twocol") if b.type is BlockType.HEADING)
    assert (heading.text, heading.level) == ("1. Introduction", 2)


@pytest.mark.parametrize("name", ["form", "twocol", "book", "math", "table"])
def test_deepseek_no_text_starts_with_a_block_marker_or_holds_html(name):
    for b in deepseek(name):
        for text in [b.text or "", *[i.text for i in b.items or []]]:
            assert not text.startswith("#"), text
            assert "<div" not in text and "<br" not in text and "<center" not in text


def test_deepseek_numbered_certification_items_stay_numbered_text():
    # The form lists "1." to "4." in separate blocks; they keep their numbers instead of
    # becoming one-item lists that Word would each number from 1.
    texts = [b.text for b in deepseek("form") if b.text]
    assert any(t.startswith("3. I am a U.S. citizen") for t in texts)


# --- headings ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ("# Title", "Title"),
        ("### Title", "Title"),
        ("## ### Title", "Title"),
        ("#Title", "#Title"),  # not a marker without a space
        ("Plain", "Plain"),
        ("## Title ##", "Title"),
    ],
)
def test_heading_text_is_stripped_and_keeps_the_label_level(text, expected):
    [out] = normalize_blocks([block(BlockType.HEADING, text, 2)])
    assert out.type is BlockType.HEADING
    assert (out.text, out.level) == (expected, 2)


def test_paragraph_that_is_a_heading_line_becomes_a_heading_with_the_marker_depth():
    [out] = normalize_blocks([block(BlockType.PARAGRAPH, "### Methods")])
    assert (out.type, out.level, out.text) == (BlockType.HEADING, 3, "Methods")


def test_heading_followed_by_body_splits_into_two_blocks():
    out = normalize_blocks([block(BlockType.PARAGRAPH, "## Results\n\nThe body text.")])
    assert [(b.type, b.text) for b in out] == [
        (BlockType.HEADING, "Results"),
        (BlockType.PARAGRAPH, "The body text."),
    ]
    assert len({b.id for b in out}) == 2


def test_paddle_vl_title_with_marker():
    [out] = vl("paragraph_title", "## 2.1 Setup")
    assert (out.type, out.text) == (BlockType.HEADING, "2.1 Setup")
    assert out.level == 3  # the section number deepens the label level


# --- other block markers ----------------------------------------------------------------


def test_bullet_lines_become_a_list_block():
    [out] = normalize_blocks([block(BlockType.PARAGRAPH, "- one\n- two\n  - nested\n* three")])
    assert out.type is BlockType.LIST
    assert [i.text for i in out.items] == ["one", "two", "three"]
    assert [c.text for c in out.items[1].children] == ["nested"]
    assert not any(i.ordered for i in out.items)


def test_numbered_lines_become_an_ordered_list_but_a_lone_number_stays_text():
    [out] = normalize_blocks([block(BlockType.PARAGRAPH, "1. first\n2. second")])
    assert out.type is BlockType.LIST and all(i.ordered for i in out.items)
    [lone] = normalize_blocks([block(BlockType.PARAGRAPH, "3. third thing")])
    assert (lone.type, lone.text) == (BlockType.PARAGRAPH, "3. third thing")


def test_quote_markers_are_dropped():
    [out] = normalize_blocks([block(BlockType.PARAGRAPH, "> quoted\n> text")])
    assert out.text == "quoted text"


def test_markdown_table_rows_become_a_table_block():
    [out] = normalize_blocks([block(BlockType.PARAGRAPH, "| a | b |\n|---|---|\n| 1 | 2 |")])
    assert out.type is BlockType.TABLE
    assert (out.table.rows, out.table.cols) == (2, 2)
    assert [c.blocks[0].text for c in out.table.cells if c.row == 1] == ["1", "2"]
    assert out.table.cells[0].header


def test_fenced_code_becomes_inline_code_without_fences():
    [out] = normalize_blocks([block(BlockType.PARAGRAPH, "```\nx = 1\n```")])
    assert out.text == "`x = 1`"


def test_display_math_block_becomes_a_formula():
    [out] = normalize_blocks([block(BlockType.PARAGRAPH, "$$\n\\frac{a}{b}\n$$")])
    assert (out.type, out.latex) == (BlockType.FORMULA, "\\frac{a}{b}")


def test_display_math_inside_text_becomes_inline_math():
    [out] = normalize_blocks([block(BlockType.PARAGRAPH, "where $$x^2$$ is large")])
    assert out.text == "where $x^2$ is large"


def test_centered_div_and_br_are_removed_from_a_caption():
    [out] = vl(
        "figure_title",
        '<div style="text-align: center;">Figure 1. The press<br>in 1850</div>',
    )
    assert (out.type, out.text) == (BlockType.CAPTION, "Figure 1. The press in 1850")


def test_html_bold_becomes_inline_markdown():
    [out] = vl("text", "A <strong>strong</strong> and <em>soft</em> word<br/>again")
    assert out.text == "A **strong** and *soft* word again"


def test_heading_marker_on_a_caption_is_dropped_but_the_type_stays():
    [out] = normalize_blocks([block(BlockType.CAPTION, "## Figure 2")])
    assert (out.type, out.text) == (BlockType.CAPTION, "Figure 2")


def test_plain_text_is_untouched_and_ids_survive():
    src = Block("p0-b7", BlockType.PARAGRAPH, BOX, 7, text="Just a sentence, 5 > 3.")
    [out] = normalize_blocks([src])
    assert (out.id, out.text) == ("p0-b7", "Just a sentence, 5 > 3.")


def test_fast_text_is_literal_and_escaped():
    [out] = blocks_from_layout(
        [LayoutItem("text", BOX, "- not a list, costs $5 and 2*3 [x]")], "fast"
    )
    assert out.type is BlockType.PARAGRAPH
    assert out.text == r"- not a list, costs \$5 and 2\*3 \[x\]"
    assert [s.text for s in parse_inline(out.text)] == ["- not a list, costs $5 and 2*3 [x]"]


def test_escape_inline_round_trips_through_the_inline_parser():
    raw = r"a_b *c* $d$ `e` [f](g) back\slash"
    assert "".join(s.text for s in parse_inline(escape_inline(raw))) == raw


# --- exporters --------------------------------------------------------------------------


def test_docx_renders_inline_markdown_without_literal_markers():
    doc = document(
        [
            block(BlockType.HEADING, "### Title", 1),
            block(
                BlockType.PARAGRAPH,
                "Plain **bold** and *italic* and ***both*** with `code`, "
                r"a [link](https://example.com), escaped \*stars\* and $x^2$ math.",
            ),
        ]
    )
    data, _ = docx_bytes(doc)
    body = etree.fromstring(zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml"))
    ns = {"w": W, "m": M}
    texts, _ = paragraphs(doc)
    assert texts[0] == "Title"
    assert "**" not in texts[1] and "`" not in texts[1] and "](" not in texts[1]
    assert "escaped *stars* and" in texts[1]
    assert len(body.findall(".//w:hyperlink", ns)) == 1
    assert len(body.findall(".//m:oMath", ns)) == 1
    runs = {
        r.findtext("w:t", namespaces=ns): (
            r.find("w:rPr/w:b", ns) is not None,
            r.find("w:rPr/w:i", ns) is not None,
        )
        for r in body.findall(".//w:r", ns)
        if r.findtext("w:t", namespaces=ns)
    }
    assert runs["bold"] == (True, False)
    assert runs["italic"] == (False, True)
    assert runs["both"] == (True, True)


def test_docx_strips_markers_that_arrive_anyway():
    doc = document(
        [
            block(BlockType.HEADING, "## ### Stacked", 2),
            block(BlockType.PARAGRAPH, "## Stray heading"),
            block(BlockType.PARAGRAPH, '<div style="text-align: center;">Centred</div>'),
            block(BlockType.CAPTION, "# Cap"),
        ]
    )
    texts, styles = paragraphs(doc)
    assert texts == ["Stacked", "Stray heading", "Centred", "Cap"]
    assert styles[0] == "Heading 2"


@pytest.mark.parametrize("name", ["form", "twocol", "book", "math", "table"])
def test_docx_of_recorded_deepseek_pages_has_no_markdown_markers(name):
    texts, _ = paragraphs(document(deepseek(name)))
    assert texts
    for text in texts:
        assert not text.startswith("#"), text
        assert "**" not in text, text
        assert "<div" not in text and "<br" not in text, text


def test_markdown_export_does_not_double_heading_markers():
    md = to_markdown(document([block(BlockType.HEADING, "## ### Title", 2)]))
    assert md.strip() == "## Title"


def test_markdown_export_keeps_literal_block_looking_text_literal():
    md = to_markdown(document([block(BlockType.PARAGRAPH, "3. third thing")]))
    assert md.strip() == r"3\. third thing"


def test_markdown_export_of_recorded_page_has_one_marker_per_heading():
    md = to_markdown(document(deepseek("form")))
    headings = [line for line in md.splitlines() if line.startswith("#")]
    assert headings
    assert all(
        line.lstrip("#").startswith(" ") and line.lstrip("# ")[0] != "#" for line in headings
    )


def test_export_helpers():
    assert strip_heading_markers("## ### x") == "x"
    assert export_text("- a", item="bullet") == "a"
    assert export_text("2. a", item="ordered") == "a"
    assert export_text("> 5 years") == "> 5 years"
    assert literal_lead("# not a heading") == r"\# not a heading"
    assert literal_lead("fine") == "fine"
