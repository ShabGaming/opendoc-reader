"""IR -> Word (.docx).

Output rules:

- Paragraphs, headings (Heading 1 to 6 by `level`), lists (bullet and numbered, nested), captions
  and footnotes use the styles of the bundled reference template (`templates/reference.docx`).
- Tables become Word tables with merged cells (gridSpan/vMerge) and a repeating header row when
  the IR marks header cells.
- Figures are inserted from their crop files, no wider than the page's text area, and keep
  with their caption.
- Display formulas are native Word equations in their own paragraph; a printed equation number
  puts the equation in a borderless two-column table with the number at the right. Inline math
  (`$...$` in text) becomes an inline equation. A formula the converter cannot handle stays as
  its LaTeX source in a monospace run, and a warning says so.
- Each source page after the first starts on a new Word page. Page furniture (header, footer,
  page number) is left out. Footnotes are moved to the end of their page's section.
- Document properties: title from the first heading (or the file name), creator "OpenDoc Reader",
  no other author information.
"""

from __future__ import annotations

import datetime as dt
import io
import re
from importlib import resources
from pathlib import Path, PureWindowsPath
from typing import Any
from urllib.parse import urlparse

from docx import Document as new_document
from docx.document import Document as DocxDocument
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Emu, Inches, Pt
from docx.table import Table as DocxTable
from docx.table import _Cell
from docx.text.paragraph import Paragraph

from opendoc_engine.blocktext import export_text
from opendoc_engine.exporters.inline import Segment, parse_inline, plain_text
from opendoc_engine.exporters.omml import (
    UnsupportedMath,
    latex_to_omath,
    split_tag,
    wrap_display,
)
from opendoc_engine.exporters.pages import OMITTED, page_blocks
from opendoc_engine.ir import Block, BlockType, Document, ListItem, Page, Table

CREATOR = "OpenDoc Reader"
_NUMBER_COLUMN = Inches(0.8)
_MARGIN = Inches(1)
_ILLEGAL_XML = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]")
_SAFE_SCHEMES = {"http", "https", "mailto"}
_HEADER_FILL = "E7E6E6"
_CHAR_WIDTH = 0.075  # inches per character, for 10 pt table text
_CELL_PADDING = 0.25  # inches of cell margin and border


def template_bytes() -> bytes:
    return (
        resources.files("opendoc_engine.exporters")
        .joinpath("templates/reference.docx")
        .read_bytes()
    )


def _safe(text: str) -> str:
    return _ILLEGAL_XML.sub("", text)


class _Sink:
    """Where blocks are written: the document body or one table cell."""

    def __init__(self, builder: _Builder, container: DocxDocument | _Cell) -> None:
        self.builder = builder
        self.container = container
        self.is_cell = isinstance(container, _Cell)
        self._fresh = self.is_cell  # a new cell already holds one empty paragraph

    def paragraph(self, style: str) -> Paragraph:
        if self._fresh:
            para = self.container.paragraphs[0]
            para.style = style
            self._fresh = False
            return para
        para = self.container.add_paragraph(style=style)
        if not self.is_cell:
            self.builder.apply_page_break(para)
        return para

    def table(self, rows: int, cols: int) -> DocxTable:
        if not self.is_cell:
            self.builder.break_before_table()
        return self.container.add_table(rows, cols)


class _Builder:
    def __init__(self, document: Document) -> None:
        self.ir = document
        self.doc: DocxDocument = new_document(io.BytesIO(template_bytes()))
        self.warnings: list[str] = []
        self.page: Page | None = None
        self.pending_break = False
        self._bullet_num = 1
        self._numbering = self.doc.part.numbering_part.element
        self.text_width = Inches(6.5)
        self.page_width = Inches(8.5)

    # --- document level -----------------------------------------------------------------

    def build(self) -> DocxDocument:
        self._setup_page()
        body = _Sink(self, self.doc)
        emitted_any = False
        for page in self.ir.pages:
            self.page = page
            if emitted_any:
                self.pending_break = True
            for block in page_blocks(page):
                before = len(self.doc.element.body)
                self.block(body, block)
                if len(self.doc.element.body) > before:
                    emitted_any = True
        self._finish()
        return self.doc

    def _setup_page(self) -> None:
        section = self.doc.sections[0]
        first = self.ir.pages[0] if self.ir.pages else None
        letter = first is None or first.height == 0 or first.width / first.height >= 0.74
        width, height = (Inches(8.5), Inches(11)) if letter else (Emu(7560310), Emu(10692130))
        if first is not None and first.width > first.height:
            width, height = height, width
        section.page_width, section.page_height = width, height
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(section, side, _MARGIN)
        self.page_width = width
        self.text_width = Emu(width - 2 * _MARGIN)

    def _finish(self) -> None:
        body = self.doc.element.body
        last = [child for child in body if child.tag != qn("w:sectPr")]
        if last and last[-1].tag == qn("w:tbl"):
            self.doc.add_paragraph(style="Normal")  # Word wants a paragraph after a final table
        self._properties()

    def _properties(self) -> None:
        title = ""
        for page in self.ir.pages:
            for block in page_blocks(page):
                if block.type is BlockType.HEADING and block.text:
                    title = plain_text(export_text(block.text, heading=True))
                    break
            if title:
                break
        if not title:
            title = PureWindowsPath(self.ir.source.path).stem
        props = self.doc.core_properties
        props.title = _safe(title)
        props.author = CREATOR
        props.last_modified_by = ""
        props.comments = ""
        props.subject = ""
        props.keywords = ""
        now = dt.datetime.now(dt.UTC).replace(microsecond=0, tzinfo=None)
        props.created = now
        props.modified = now
        props.revision = 1

    def apply_page_break(self, para: Paragraph) -> None:
        if self.pending_break:
            para.paragraph_format.page_break_before = True
            self.pending_break = False

    def break_before_table(self) -> None:
        if self.pending_break:  # a table cannot carry the break itself, so use a thin paragraph
            spacer = self.doc.add_paragraph(style="Normal")
            spacer.paragraph_format.space_after = Pt(0)
            spacer.paragraph_format.line_spacing = Pt(1)
            self.apply_page_break(spacer)

    def warn(self, message: str) -> None:
        if message not in self.warnings:
            self.warnings.append(message)

    # --- blocks ---------------------------------------------------------------------------

    def block(self, sink: _Sink, block: Block) -> None:
        kind = block.type
        if kind is BlockType.HEADING:
            level = min(max(block.level or 1, 1), 6)
            self._text_block(sink, block, f"Heading {level}")
        elif kind is BlockType.CAPTION:
            self._text_block(sink, block, "Caption")
        elif kind is BlockType.FOOTNOTE:
            self._text_block(sink, block, "Footnote Text")
        elif kind is BlockType.LIST:
            self._list(sink, block)
        elif kind is BlockType.TABLE:
            self._table(sink, block)
        elif kind is BlockType.FORMULA:
            self._formula(sink, block)
        elif kind is BlockType.FIGURE:
            self._figure(sink, block)
        elif kind not in OMITTED:
            self._text_block(sink, block, "Table Text" if sink.is_cell else "Normal")

    def _text_block(self, sink: _Sink, block: Block, style: str) -> None:
        if not (block.text and block.text.strip()):
            return
        if sink.is_cell and style == "Normal":
            style = "Table Text"
        text = export_text(block.text, heading=block.type is BlockType.HEADING)
        if not text:
            return
        para = sink.paragraph(style)
        self.inline(para, text, block.id)

    # --- inline content -----------------------------------------------------------------

    def inline(self, para: Paragraph, text: str, block_id: str) -> None:
        for seg in parse_inline(_safe(text)):
            if seg.kind == "math":
                self._inline_math(para, seg, block_id)
            elif seg.kind == "link":
                self._link(para, seg)
            else:
                self._run(para, seg.text, seg, code=seg.kind == "code")

    def _run(self, para: Paragraph, text: str, seg: Segment, *, code: bool = False) -> None:
        run = para.add_run(text)
        run.bold = seg.bold or None
        run.italic = seg.italic or None
        if code:
            run.style = "Code"

    def _link(self, para: Paragraph, seg: Segment) -> None:
        if urlparse(seg.url).scheme.lower() not in _SAFE_SCHEMES:
            self._run(para, seg.text, seg)
            return
        rid = para.part.relate_to(seg.url, RT.HYPERLINK, is_external=True)
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), rid)
        link.set(qn("w:history"), "1")
        run = parse_xml(
            f'<w:r {nsdecls("w")}><w:rPr><w:rStyle w:val="Hyperlink"/></w:rPr>'
            '<w:t xml:space="preserve"></w:t></w:r>'
        )
        run.find(qn("w:t")).text = seg.text
        link.append(run)
        para._p.append(link)

    def _inline_math(self, para: Paragraph, seg: Segment, block_id: str) -> None:
        try:
            para._p.append(latex_to_omath(seg.text, display=False))
        except UnsupportedMath as exc:
            self.warn(f"inline math in {block_id} kept as LaTeX text ({exc})")
            self._run(para, f"${seg.text}$", seg, code=True)

    # --- lists ------------------------------------------------------------------------------

    def _list(self, sink: _Sink, block: Block) -> None:
        if not block.items:
            return
        decimal_num = self._new_decimal_num()
        self._list_items(sink, block, block.items, 0, decimal_num)

    def _list_items(
        self, sink: _Sink, block: Block, items: list[ListItem], depth: int, decimal_num: int
    ) -> None:
        for item in items:
            para = sink.paragraph("List Paragraph")
            props = para._p.get_or_add_pPr()
            num = props.get_or_add_numPr()
            num.get_or_add_ilvl().val = min(depth, 8)
            num.get_or_add_numId().val = decimal_num if item.ordered else self._bullet_num
            if sink.is_cell:
                para.paragraph_format.space_after = Pt(2)
            self.inline(
                para, export_text(item.text, item="ordered" if item.ordered else "bullet"), block.id
            )
            self._list_items(sink, block, item.children, depth + 1, decimal_num)

    def _new_decimal_num(self) -> int:
        """A fresh list instance of the decimal definition, so each list starts at 1."""
        numbers = [int(n.get(qn("w:numId"))) for n in self._numbering.findall(qn("w:num"))]
        num_id = max(numbers, default=0) + 1
        overrides = "".join(
            f'<w:lvlOverride w:ilvl="{i}"><w:startOverride w:val="1"/></w:lvlOverride>'
            for i in range(3)
        )
        self._numbering.append(
            parse_xml(
                f'<w:num {nsdecls("w")} w:numId="{num_id}"><w:abstractNumId w:val="2"/>'
                f"{overrides}</w:num>"
            )
        )
        return num_id

    # --- tables -----------------------------------------------------------------------------

    def _table(self, sink: _Sink, block: Block) -> None:
        table = block.table
        if table is None or table.rows == 0 or table.cols == 0:
            return
        grid = sink.table(table.rows, table.cols)
        grid.style = "Table Grid"
        merged = self._merge_cells(grid, table, block.id)
        self._set_widths(grid, table, None if sink.is_cell else self.text_width)
        header_rows = _header_rows(table)
        for cell in sorted(table.cells, key=lambda c: (c.row, c.col)):
            if (cell.row, cell.col) not in merged:
                continue
            target = grid.cell(cell.row, cell.col)
            inner = _Sink(self, target)
            for sub in cell.blocks:
                self.block(inner, sub)
            if cell.row in header_rows:
                _shade(target)
                for para in target.paragraphs:
                    for run in para.runs:
                        run.bold = True
        for index, row in enumerate(grid.rows):
            props = row._tr.get_or_add_trPr()
            props.append(OxmlElement("w:cantSplit"))
            if index in header_rows:
                props.append(OxmlElement("w:tblHeader"))

    def _set_widths(self, grid: DocxTable, table: Table, total: int | None) -> None:
        """Fixed column widths: every column gets room for its longest word, and what is left
        goes to the columns that have longer text to wrap."""
        if total is None:
            return
        minimum = [0.0] * table.cols
        preferred = [0.0] * table.cols
        for cell in table.cells:
            if cell.col_span != 1 or cell.col >= table.cols:
                continue
            texts = [line for b in cell.blocks for line in (b.text or "").split("\n")]
            word = max((len(w) for t in texts for w in t.split()), default=0)
            line = max((len(t) for t in texts), default=0)
            minimum[cell.col] = max(minimum[cell.col], _CELL_PADDING + _CHAR_WIDTH * word)
            preferred[cell.col] = max(preferred[cell.col], _CELL_PADDING + _CHAR_WIDTH * line)
        minimum = [max(m, _CELL_PADDING + 3 * _CHAR_WIDTH) for m in minimum]
        preferred = [max(p, m) for p, m in zip(preferred, minimum, strict=True)]
        inches = total / 914400
        if sum(preferred) <= inches:
            spare = inches - sum(preferred)
            widths_in = [p + spare * p / sum(preferred) for p in preferred]
        elif sum(minimum) >= inches:
            widths_in = [m * inches / sum(minimum) for m in minimum]
        else:
            room = (inches - sum(minimum)) / sum(
                p - m for p, m in zip(preferred, minimum, strict=True)
            )
            widths_in = [m + (p - m) * room for p, m in zip(preferred, minimum, strict=True)]
        widths = [Emu(int(w * 914400)) for w in widths_in]
        grid.autofit = False
        for column, width in zip(grid.columns, widths, strict=True):
            column.width = width
        for row in grid.rows:
            col = 0
            for tc in row._tr.tc_lst:
                span = tc.grid_span
                tc.width = Emu(sum(widths[col : col + span]))
                col += span

    def _merge_cells(self, grid: DocxTable, table: Table, block_id: str) -> set[tuple[int, int]]:
        """Merge spanning cells. Returns the anchors that were kept (others were dropped)."""
        occupied: set[tuple[int, int]] = set()
        anchors: set[tuple[int, int]] = set()
        for cell in sorted(table.cells, key=lambda c: (c.row, c.col)):
            r2 = min(cell.row + max(cell.row_span, 1), table.rows) - 1
            c2 = min(cell.col + max(cell.col_span, 1), table.cols) - 1
            slots = {(r, c) for r in range(cell.row, r2 + 1) for c in range(cell.col, c2 + 1)}
            if cell.row >= table.rows or cell.col >= table.cols or slots & occupied:
                self.warn(f"table {block_id}: cell at row {cell.row}, column {cell.col} dropped")
                continue
            occupied |= slots
            anchors.add((cell.row, cell.col))
            if len(slots) > 1:
                grid.cell(cell.row, cell.col).merge(grid.cell(r2, c2))
        return anchors

    # --- formulas -------------------------------------------------------------------------

    def _formula(self, sink: _Sink, block: Block) -> None:
        source = block.latex if block.latex is not None else (block.text or "")
        latex, tag = split_tag(_safe(source))
        number = block.number or tag
        if not latex.strip():
            return
        omath = None
        try:
            omath = latex_to_omath(latex, display=True)
        except UnsupportedMath as exc:
            self.warn(f"formula {block.id} kept as LaTeX text (unsupported construct: {exc})")

        if number and not sink.is_cell:
            self._numbered_formula(sink, latex, omath, _bracketed(number))
            return
        para = sink.paragraph("Equation")
        self._put_formula(para, latex, omath)

    def _put_formula(self, para: Paragraph, latex: str, omath: Any) -> None:
        if omath is not None:
            para._p.append(wrap_display(omath))
        else:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = para.add_run(latex.strip())
            run.style = "Code"

    def _numbered_formula(self, sink: _Sink, latex: str, omath: Any, number: str) -> None:
        grid = sink.table(1, 2)
        grid.autofit = False
        widths = [Emu(self.text_width - _NUMBER_COLUMN), _NUMBER_COLUMN]
        for column, width in zip(grid.columns, widths, strict=True):
            column.width = width
        for cell, width in zip(grid.rows[0].cells, widths, strict=True):
            cell.width = width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        math_para = grid.cell(0, 0).paragraphs[0]
        math_para.style = "Equation"
        self._put_formula(math_para, latex, omath)
        number_para = grid.cell(0, 1).paragraphs[0]
        number_para.style = "Equation"
        number_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        number_para.add_run(_safe(number))

    # --- figures ----------------------------------------------------------------------------

    def _figure(self, sink: _Sink, block: Block) -> None:
        crop = Path(block.figure_crop) if block.figure_crop else None
        if crop is None or not crop.is_file():
            self.warn(f"figure {block.id} skipped: crop image not found")
            return
        width = self._figure_width(block)
        para = sink.paragraph("Table Text" if sink.is_cell else "Normal")
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        try:
            para.add_run().add_picture(str(crop), width=width)
        except Exception as exc:  # python-docx raises on image formats it cannot read
            self.warn(f"figure {block.id} skipped: image not readable ({type(exc).__name__})")
            para._p.getparent().remove(para._p)
            return
        if not sink.is_cell:
            para.paragraph_format.keep_with_next = True

    def _figure_width(self, block: Block) -> Emu:
        page = self.page
        full = self.text_width
        if page is None or page.width <= 0:
            return full
        fraction = (block.bbox[2] - block.bbox[0]) / page.width
        if fraction <= 0:
            return full
        return Emu(int(min(full, fraction * self.page_width)))


def _bracketed(number: str) -> str:
    number = number.strip()
    return number if number.startswith("(") else f"({number})"


def _header_rows(table: Table) -> set[int]:
    """Leading rows in which at least one cell is marked as a header."""
    rows: set[int] = set()
    for row in range(table.rows):
        if any(c.header for c in table.cells if c.row == row):
            rows.add(row)
        else:
            break
    return rows


def _shade(cell: _Cell) -> None:
    props = cell._tc.get_or_add_tcPr()
    props.append(
        parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{_HEADER_FILL}"/>')
    )


def build_docx(document: Document) -> tuple[DocxDocument, list[str]]:
    """Render `document` to a python-docx document. Returns it with the list of warnings."""
    builder = _Builder(document)
    doc = builder.build()
    return doc, builder.warnings


def docx_bytes(document: Document) -> tuple[bytes, list[str]]:
    doc, warnings = build_docx(document)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue(), warnings
