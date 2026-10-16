"""Developer check: run the exporters on real recognition output from the model spike.

Maps the PaddleOCR-VL result JSON of the spike pages (book, table, math) to IR with a simple
label mapping, writes them through the real RPC methods (`document.assemble`, `export.docx`,
`export.markdown`) and leaves the results in an output folder for opening in Word. Figure crops
are cut from the page images. It reads from the spike folder and never writes there.

    uv run --with pillow python scripts/export_spike_samples.py [--out G:\\opendoc-dev]

Needs Pillow only for the figure crops (the engine itself gets it from the pipeline lane).
"""

from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
from pathlib import Path

from opendoc_engine.htmltable import parse_html_table
from opendoc_engine.ir import Block, BlockType, Page
from opendoc_engine.methods import build_server

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
SPIKE = Path(r"G:\opendoc-spike\results\bench")
PVL = SPIKE / "pvl-cuda"
DS1_MATH = SPIKE / "ds1-q8-cuda" / "math.md"

LABELS = {
    "header": BlockType.HEADER,
    "header_image": BlockType.HEADER,
    "footer": BlockType.FOOTER,
    "footer_image": BlockType.FOOTER,
    "number": BlockType.PAGE_NUMBER,
    "text": BlockType.PARAGRAPH,
    "abstract": BlockType.PARAGRAPH,
    "content": BlockType.PARAGRAPH,
    "figure_title": BlockType.CAPTION,
    "footnote": BlockType.FOOTNOTE,
    "aside_text": BlockType.OTHER,
}
HEADINGS = {"doc_title": 1, "paragraph_title": 2}
FIGURES = {"image", "chart", "seal"}


def _latex(content: str) -> str:
    match = re.search(r"\$\$(.*)\$\$", content, re.DOTALL)
    return (match.group(1) if match else content).strip()


def _paddle_page(name: str, index: int, work: Path) -> Page:
    from PIL import Image

    data = json.loads((PVL / f"{name}.json").read_text(encoding="utf-8"))
    source = Path(data["input_path"])
    image = Image.open(source)
    blocks: list[Block] = []
    numbers: list[tuple[str, tuple[float, float, float, float]]] = []
    items = sorted(
        data["parsing_res_list"],
        key=lambda b: (b["block_order"] is None, b["block_order"] or 0, b["block_id"]),
    )
    for order, raw in enumerate(items):
        label = raw["block_label"]
        bbox = tuple(float(v) for v in raw["block_bbox"])
        block_id = f"p{index}-b{raw['block_id']}"
        content = raw.get("block_content") or ""
        if label == "formula_number":
            numbers.append((content.strip().strip("()"), bbox))
            continue
        if label == "display_formula":
            blocks.append(Block(block_id, BlockType.FORMULA, bbox, order, latex=_latex(content)))
        elif label == "table":
            table = parse_html_table(content, bbox, block_id)
            if table:
                blocks.append(Block(block_id, BlockType.TABLE, bbox, order, table=table))
        elif label in FIGURES:
            crop = work / "crops" / f"{block_id}.png"
            crop.parent.mkdir(parents=True, exist_ok=True)
            image.crop(tuple(int(v) for v in bbox)).save(crop)
            blocks.append(Block(block_id, BlockType.FIGURE, bbox, order, figure_crop=str(crop)))
        elif label in HEADINGS:
            blocks.append(
                Block(block_id, BlockType.HEADING, bbox, order, level=HEADINGS[label], text=content)
            )
        else:
            kind = LABELS.get(label, BlockType.PARAGRAPH)
            blocks.append(Block(block_id, kind, bbox, order, text=content))

    formulas = [b for b in blocks if b.type is BlockType.FORMULA]
    for number, nbox in numbers:  # attach each number to the formula it overlaps vertically
        best, overlap = None, 0.0
        for formula in formulas:
            top, bottom = max(nbox[1], formula.bbox[1]), min(nbox[3], formula.bbox[3])
            if bottom - top > overlap and formula.number is None:
                best, overlap = formula, bottom - top
        if best is not None:
            best.number = number
    return Page(index, data["width"], data["height"], str(source), blocks)


def _ds1_math_page(work: Path) -> Page:
    """DeepSeek-OCR markdown of the math page: `\\[..\\]` is display math, `\\(..\\)` inline."""
    text = DS1_MATH.read_text(encoding="utf-8")
    blocks: list[Block] = []
    box = (0.0, 0.0, 1376.0, 2200.0)

    def add(kind: BlockType, **fields) -> None:
        blocks.append(Block(f"ds-b{len(blocks)}", kind, box, len(blocks), **fields))

    for chunk in re.split(r"\n\s*\n", text):
        chunk = chunk.strip()
        if not chunk:
            continue
        displays = re.findall(r"\\\[(.*?)\\\]", chunk, re.DOTALL)
        leading = re.match(r"\((\d+)\)\s*\\\(", chunk)
        if displays and re.fullmatch(r"(\s*\\\[.*?\\\]\s*)+", chunk, re.DOTALL):
            for body in displays:
                body = body.strip()
                number = None
                tag = re.match(r"\(\d+\)\s*\\qquad\s*(.*)", body, re.DOTALL)
                if tag:
                    number, body = re.match(r"\((\d+)\)", body).group(1), tag.group(1)
                if body.startswith("{") and body.endswith("}"):
                    body = body[1:-1]
                add(BlockType.FORMULA, latex=body, number=number)
        elif leading and chunk.endswith("\\)"):
            add(
                BlockType.FORMULA, latex=chunk[chunk.index("\\(") + 2 : -2], number=leading.group(1)
            )
        else:
            add(BlockType.PARAGRAPH, text=re.sub(r"\\\((.*?)\\\)", r"$\1$", chunk, flags=re.DOTALL))
    return Page(0, 1376, 2200, "ds1-math.png", blocks)


def _call(server, method: str, params: dict, call_id: int) -> dict:
    line = json.dumps({"jsonrpc": "2.0", "id": call_id, "method": method, "params": params})
    reply = json.loads(server.handle_line(line))
    if "error" in reply:
        raise SystemExit(f"{method} failed: {reply['error']}")
    return reply["result"]


def _run(server, out: Path, name: str, pages: list[Page]) -> None:
    work = out / "work" / name
    exports = out / "exports"
    shutil.rmtree(work / "ir", ignore_errors=True)
    (work / "ir").mkdir(parents=True)
    page_files = []
    for i, page in enumerate(pages):
        path = work / "ir" / f"page-{i:04d}.json"
        path.write_text(json.dumps(page.to_dict(), ensure_ascii=False), encoding="utf-8")
        page_files.append(str(path))
    exports.mkdir(parents=True, exist_ok=True)
    docx_path = exports / f"{name}.docx"
    md_path = exports / f"{name}.md"
    assets = exports / f"{name}_assets"
    docx_path.unlink(missing_ok=True)
    md_path.unlink(missing_ok=True)
    shutil.rmtree(assets, ignore_errors=True)

    document = work / "document.json"
    _call(
        server,
        "document.assemble",
        {
            "source": {"path": f"{name}.pdf", "kind": "pdf", "pageCount": len(pages)},
            "engine": {"id": "paddleocr-vl", "model": "spike", "version": "0.1.0"},
            "pages": page_files,
            "outPath": str(document),
        },
        1,
    )
    docx = _call(
        server, "export.docx", {"document": str(document), "outputPath": str(docx_path)}, 2
    )
    md = _call(
        server,
        "export.markdown",
        {"document": str(document), "outputPath": str(md_path), "assetsDir": str(assets)},
        3,
    )
    print(f"{name}: {docx_path.name}, {md_path.name}, {len(md['files']) - 1} asset(s)")
    for warning in docx["warnings"] + md["warnings"]:
        print(f"  warning: {warning}")
    print("  " + _verify(docx_path))


def _verify(path: Path) -> str:
    """Open the .docx with python-docx and check every equation against the OMML validator."""
    import docx
    from lxml import etree

    from opendoc_engine.exporters.omml import M_NS
    from opendoc_engine.exporters.omml_check import check_omml

    package = docx.Document(str(path))
    body = package.element.body
    find = etree.XPath(".//m:oMathPara | .//w:p/m:oMath", namespaces={"m": M_NS, "w": W_NS})
    math = find(body)
    problems = [e for node in math for e in check_omml(node)]
    if problems:
        raise SystemExit(f"{path.name}: invalid OMML: {problems[:3]}")
    return (
        f"python-docx opened it: {len(package.paragraphs)} paragraphs, "
        f"{len(package.tables)} tables, {len(math)} equations (valid), "
        f"{len(package.inline_shapes)} pictures"
    )


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--out", type=Path, default=Path(r"G:\opendoc-dev"))
    args = parser.parse_args()
    if SPIKE in args.out.parents or args.out == SPIKE:
        raise SystemExit("refusing to write into the spike folder")

    server = build_server()
    samples = {}
    for index, name in enumerate(["book", "table", "math"]):
        samples[name] = _paddle_page(name, index, args.out / "work" / name)
    for name, page in samples.items():
        _run(server, args.out, f"spike-{name}", [page])
    combined = [
        Page(i, p.width, p.height, p.source_image, p.blocks) for i, p in enumerate(samples.values())
    ]
    _run(server, args.out, "spike-3pages", combined)
    _run(server, args.out, "spike-math-deepseek", [_ds1_math_page(args.out / "work")])
    return 0


if __name__ == "__main__":
    sys.exit(main())
