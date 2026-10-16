"""Helper for the end-to-end test (`tests/e2e.rs`), run with a Python that has the engine's
dependencies and its dev group (the repository's `engine/.venv`).

    e2e_helper.py make-inputs OUT_DIR SAMPLES_DIR
    e2e_helper.py check-docx FILE.docx
    e2e_helper.py check-ir PAGE_IR.json...

Every command prints one JSON object and exits non-zero when a check fails.
"""

from __future__ import annotations

import json
import shutil
import sys
from pathlib import Path

ENGINE = Path(__file__).resolve().parents[5] / "engine"


def make_inputs(out: Path, samples: Path) -> dict:
    from PIL import Image

    out.mkdir(parents=True, exist_ok=True)
    pages = [
        samples / "ia-tom-sawyer-1876" / "n100_2200.png",
        samples / "ia-statabstract-1920" / "n120_2200.png",
        samples / "ia-granville-calculus-1911" / "n100_2200.png",
    ]
    images = [Image.open(p).convert("RGB") for p in pages]
    scanned = out / "scanned-three-pages.pdf"
    images[0].save(scanned, "PDF", save_all=True, append_images=images[1:], resolution=200.0)

    single = out / "form-page.png"
    shutil.copyfile(samples / "irs-fw9" / "page01_200dpi.png", single)

    sys.path.insert(0, str(ENGINE / "tests"))
    from pdfgen import PageSpec, Text, make_pdf

    born = out / "born-digital.pdf"
    born.write_bytes(
        make_pdf(
            [
                PageSpec(
                    [
                        Text(72, 700, 22, "Quarterly report", bold=True),
                        Text(72, 660, 12, "Revenue grew in the third quarter."),
                        Text(72, 640, 12, "Costs stayed flat over the same period."),
                    ]
                ),
                PageSpec(
                    [
                        Text(72, 700, 16, "Outlook", bold=True),
                        Text(72, 670, 12, "We expect steady growth next year."),
                    ]
                ),
            ]
        )
    )
    return {"files": [str(scanned), str(single), str(born)]}


def check_docx(path: Path) -> dict:
    from docx import Document

    document = Document(str(path))
    text = "\n".join(p.text for p in document.paragraphs)
    result = {
        "paragraphs": len(document.paragraphs),
        "tables": len(document.tables),
        "images": len(document.inline_shapes),
        "chars": len(text),
    }
    if result["paragraphs"] == 0 and result["tables"] == 0:
        raise SystemExit(json.dumps({"error": "empty document", **result}))
    return result


def check_ir(paths: list[Path]) -> dict:
    from jsonschema import Draft202012Validator

    schema = json.loads((ENGINE / "schema" / "ir.schema.json").read_text(encoding="utf-8"))
    validator = Draft202012Validator({"$ref": "#/$defs/Page", "$defs": schema["$defs"]})
    problems = {}
    blocks = 0
    for path in paths:
        page = json.loads(path.read_text(encoding="utf-8"))
        errors = [e.message for e in validator.iter_errors(page)]
        if not Path(page["sourceImage"]).is_file():
            errors.append("sourceImage does not exist")
        blocks += len(page["blocks"])
        if errors:
            problems[str(path)] = errors[:5]
    if problems:
        raise SystemExit(json.dumps({"error": "invalid IR", "problems": problems}))
    return {"pages": len(paths), "blocks": blocks}


def main() -> None:
    command, args = sys.argv[1], sys.argv[2:]
    if command == "make-inputs":
        result = make_inputs(Path(args[0]), Path(args[1]))
    elif command == "check-docx":
        result = check_docx(Path(args[0]))
    elif command == "check-ir":
        result = check_ir([Path(a) for a in args])
    else:
        raise SystemExit(f"unknown command {command}")
    print(json.dumps(result))


if __name__ == "__main__":
    main()
